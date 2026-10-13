"""Private Excel and CSV exports (spec 12.1, 12.2 Export)."""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from collector.config import Settings
from collector.issues import publish
from collector.pipeline import process_raw

SHEETS = ["Summary", "Top & Bottom 10", "Models", "Engines", "Payment plans"]


@pytest.fixture
def exported(settings: Settings, synthetic: list[dict[str, Any]]) -> Settings:
    for raw in synthetic:
        process_raw(settings, raw)
    return settings


def test_workbook_sheets_formulas_no_identifiers(exported: Settings) -> None:
    wb = load_workbook(exported.data / "exports" / "2026-10-13.xlsx")
    assert wb.sheetnames == SHEETS
    headers = [c.value for ws in wb for c in ws[1] if c.value]
    assert not [h for h in headers if str(h).lower() in ("vin", "offer link", "link", "phone", "telefon")]
    formulas = [c.value for ws in wb for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert formulas and not [f for f in formulas if "#REF!" in f]
    total = [r for r in wb["Models"].iter_rows(values_only=True) if r[0] == "Total"][0]
    assert str(total[2]).startswith("=SUM(C2:C")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_recalculates_without_errors(exported: Settings, tmp_path: Path) -> None:
    src = exported.data / "exports" / "2026-10-13.xlsx"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path), str(src)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmp_path / src.name, data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value in ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A")]
    assert errors == []
    assert wb["Summary"]["B2"].value == 116


def test_latest_csv(exported: Settings) -> None:
    with (exported.data / "exports" / "latest.csv").open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0][:3] == ["Model", "Engine", "Model code"]
    assert "VIN" not in rows[0] and "Offer link" not in rows[0]
    assert len(rows) == 116 + 1
    special = rows[0].index("Special price (PLN)")
    assert all(r[special] == "" or r[special].isdigit() for r in rows[1:])


def test_issue_dry_run(exported: Settings, capsys: pytest.CaptureFixture[str]) -> None:
    p = exported.data.parent / ".run" / "result.json"
    res = json.loads(p.read_text(encoding="utf-8"))
    res["issues"] = [{"label": "publication-guard", "title": "t", "body": "b"}]
    p.write_text(json.dumps(res), encoding="utf-8")
    assert publish(p, dry_run=True) == ["created (publication-guard)"]
    assert "gh issue list --label publication-guard --state open" in capsys.readouterr().out
