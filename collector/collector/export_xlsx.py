"""Per-run Excel workbook and latest.csv, private repository only (spec 4.2 step 10, 7.3 offer columns).

Sheets follow the 13 Oct 2026 workbook: Summary, Top & Bottom 10, Models, Engines, Payment plans, without VIN,
link or phone columns (spec 16.1). Totals and ratios are formulas; output bytes are deterministic.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from collector.rules import PROMO_KEYS, Rules

CSV_COLUMNS = [
    ("model", "Model"), ("engine", "Engine"), ("code", "Model code"), ("model_year", "Model year"),
    ("catalog", "Catalog price (PLN)"), ("special", "Special price (PLN)"), ("discount", "Discount (PLN)"),
    ("discount_pct", "Discount %"), ("price_30d", "30-day price (PLN)"), ("promos", "Promos"),
    ("rate", "Rate (PLN net)"), ("plan_code", "Plan"), ("initial_fee", "Initial fee (PLN)"),
    ("term", "Term (months)"), ("identical", "Identical offers"), ("dealer", "Dealer"), ("score", "Deal score"),
    ("offer_id", "Offer"),
]
PLN = '#,##0 "zł"'
PCT = "0.0%"


def _promos(bits: int, rules: Rules) -> str:
    return "; ".join(rules.promo["promos"][k]["name"] for k in PROMO_KEYS
                     if bits >> int(rules.promo["promos"][k]["badge_bit"]) & 1)


def write_csv(path: Path, rows: list[dict[str, Any]], rules: Rules) -> None:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([label for _, label in CSV_COLUMNS])
    for r in rows:
        vals = []
        for key, _ in CSV_COLUMNS:
            v = _promos(r["promo_bits"], rules) if key == "promos" else r.get(key)
            vals.append("" if v is None else v)
        w.writerow(vals)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = buf.getvalue().encode("utf-8-sig")
    if not (path.exists() and path.read_bytes() == data):
        path.write_bytes(data)


def _header(ws: Any, labels: list[str]) -> None:
    ws.append(labels)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _widths(ws: Any, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_xlsx(path: Path, date: str, rows: list[dict[str, Any]], rules: Rules,
               directory: list[dict[str, Any]]) -> None:
    wb = Workbook()
    fixed = datetime.fromisoformat(date + "T00:00:00")
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "audilitics collector"
    wb.properties.lastModifiedBy = "audilitics collector"

    # ---- Models ------------------------------------------------------------------------------------------
    ws_m = wb.active
    ws_m.title = "Models"
    by_model: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        by_model.setdefault(r["model"], []).append(r)
    labels = ["Model", "Segment", "Offers", "Min special (PLN)", "Median special (PLN)", "Max special (PLN)",
              "Median discount %", "Min rate (PLN net)", *[rules.promo["promos"][k]["name"] for k in PROMO_KEYS]]
    _header(ws_m, labels)
    model_rows = [m for m in directory if by_model.get(m["name"])]
    for m in model_rows:
        cars = by_model[m["name"]]
        sp = [c["special"] for c in cars if c["special"] is not None]
        dp = [c["discount_pct"] for c in cars if c["discount_pct"] is not None]
        rt = [c["rate"] for c in cars if c["rate"] is not None]
        ws_m.append([m["name"], m["segment"], len(cars), min(sp) if sp else None, median(sp) if sp else None,
                     max(sp) if sp else None, median(dp) if dp else None, min(rt) if rt else None,
                     *[sum(1 for c in cars if c["promo_bits"] >> int(rules.promo["promos"][k]["badge_bit"]) & 1)
                       for k in PROMO_KEYS]])
    last = ws_m.max_row
    total_row = last + 1
    ws_m.append(["Total", None, f"=SUM(C2:C{last})", f"=MIN(D2:D{last})", None, f"=MAX(F2:F{last})", None,
                 f"=MIN(H2:H{last})", *[f"=SUM({get_column_letter(9 + i)}2:{get_column_letter(9 + i)}{last})"
                                       for i in range(len(PROMO_KEYS))]])
    for c in ws_m[total_row]:
        c.font = Font(bold=True)
    for row in ws_m.iter_rows(min_row=2, max_row=total_row):
        for idx in (3, 4, 5, 7):
            row[idx].number_format = PLN
        row[6].number_format = PCT
    _widths(ws_m, [30, 9, 7, 16, 18, 16, 16, 16, *[14] * len(PROMO_KEYS)])

    # ---- Summary (first sheet) -----------------------------------------------------------------------------
    ws_s = wb.create_sheet("Summary", 0)
    priced = [r for r in rows if r["special"] is not None]
    ws_s.append(["Audilitics snapshot", date])
    ws_s.append(["Offers in stock", f"=Models!C{total_row}"])
    ws_s.append(["Models with offers", f"=COUNTA(Models!A2:A{last})"])
    ws_s.append(["Dealers with offers", len({r["dealer_nr"] for r in rows})])
    ws_s.append(["Configurations", len({r["config_key"] for r in rows})])
    ws_s.append(["Lowest special price (PLN)", f"=Models!D{total_row}"])
    ws_s.append(["Highest special price (PLN)", f"=Models!F{total_row}"])
    ws_s.append(["Median discount %", median([r["discount_pct"] for r in priced]) if priced else None])
    ws_s.append(["Lowest rate (PLN net)", f"=Models!H{total_row}"])
    ws_s.append(["Offers without a special price", len(rows) - len(priced)])
    for i, k in enumerate(PROMO_KEYS):
        col = get_column_letter(9 + i)
        ws_s.append([f"Badge: {rules.promo['promos'][k]['name']}", f"=Models!{col}{total_row}"])
    ws_s.append(["Source", "gotowedoodbioru.audi.pl stock feed, redacted at ingestion; private, not for publication"])
    for c in ws_s["A"]:
        c.font = Font(bold=True)
    for ref in ("B6", "B7"):
        ws_s[ref].number_format = PLN
    ws_s["B8"].number_format = PCT
    _widths(ws_s, [44, 40])

    # ---- Top & Bottom 10 -----------------------------------------------------------------------------------
    ws_t = wb.create_sheet("Top & Bottom 10", 1)
    scored = sorted([r for r in rows if r["score"] is not None],
                    key=lambda r: (-r["score"], -(r["discount"] or 0), r["offer_id"]))
    _header(ws_t, ["Rank", "Group", "Model", "Engine", "Catalog price (PLN)", "Special price (PLN)", "Discount %",
                   "Deal score", "Rate (PLN net)", "Plan", "Dealer", "Offer"])
    groups = [("Top 10", scored[:10]), ("Bottom 10", list(reversed(scored[-10:])) if len(scored) > 10 else [])]
    for gname, lst in groups:
        for i, r in enumerate(lst, start=1):
            n = ws_t.max_row + 1
            ws_t.append([i, gname, r["model"], r["engine"], r["catalog"], r["special"], f"=1-F{n}/E{n}",
                         r["score"], r["rate"], r["plan_code"], r["dealer"], r["offer_id"]])
            for col in "EFI":
                ws_t[f"{col}{n}"].number_format = PLN
            ws_t[f"G{n}"].number_format = PCT
    _widths(ws_t, [6, 10, 28, 36, 16, 16, 11, 10, 13, 9, 30, 16])

    # ---- Engines -------------------------------------------------------------------------------------------
    ws_e = wb.create_sheet("Engines")
    _header(ws_e, ["Model", "Engine group", "Offers", "Min special (PLN)", "Median special (PLN)",
                   "Min rate (PLN net)", "Share of model"])
    for m in model_rows:
        cars = by_model[m["name"]]
        groups_e: dict[str, list[dict[str, Any]]] = {}
        for c in cars:
            groups_e.setdefault(c["engine_group"] or "(none)", []).append(c)
        first = ws_e.max_row + 1
        lastm = first + len(groups_e) - 1
        for g in sorted(groups_e):
            cs = groups_e[g]
            sp = [c["special"] for c in cs if c["special"] is not None]
            rt = [c["rate"] for c in cs if c["rate"] is not None]
            n = ws_e.max_row + 1
            ws_e.append([m["name"], g, len(cs), min(sp) if sp else None, median(sp) if sp else None,
                         min(rt) if rt else None, f"=C{n}/SUM(C{first}:C{lastm})"])
            for col in "DEF":
                ws_e[f"{col}{n}"].number_format = PLN
            ws_e[f"G{n}"].number_format = PCT
    _widths(ws_e, [30, 18, 7, 16, 18, 16, 14])

    # ---- Payment plans -------------------------------------------------------------------------------------
    ws_p = wb.create_sheet("Payment plans")
    _header(ws_p, ["Plan", "Footnote", "Initial fee (of catalog price)", "Term (months)", "Mileage per year (km)",
                   "Offers", "Share of offers", "Lowest rate (PLN net)"])
    plans = rules.plans["plans"]
    first = 2
    lastp = first + len(plans) - 1
    for code, p in plans.items():
        cs = [r for r in rows if r["plan_code"] == code]
        rt = [r["rate"] for r in cs if r["rate"] is not None]
        n = ws_p.max_row + 1
        ws_p.append([code, p["footnote"], p["fee_pct"], p["term_months"], p["mileage_km"], len(cs),
                     f"=F{n}/SUM(F{first}:F{lastp})", min(rt) if rt else None])
        ws_p[f"C{n}"].number_format = "0%"
        ws_p[f"G{n}"].number_format = PCT
        ws_p[f"H{n}"].number_format = PLN
    l101 = rules.plans.get("leasing_101") or {}
    ws_p.append([])
    ws_p.append(["Leasing 101% (flag, not a rate)", None, l101.get("fee_pct"), l101.get("term_months"), None, None,
                 None, None])
    ws_p.append(["Rates are Audi Perfect Lease, PLN net, for businesses, calculated from the catalog price."])
    _widths(ws_p, [34, 9, 26, 13, 20, 8, 15, 20])

    buf = io.BytesIO()
    wb.save(buf)
    _write_deterministic_zip(path, buf.getvalue(), date)


def _write_deterministic_zip(path: Path, data: bytes, date: str) -> None:
    """Re-pack the xlsx with fixed member timestamps so equal content gives equal bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps dcterms:modified with the wall clock at save time
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)",
                              rb"\g<1>" + f"{date}T00:00:00Z".encode() + rb"\g<2>", body)
            dst.writestr(zi, body)
    blob = out.getvalue()
    path.parent.mkdir(parents=True, exist_ok=True)
    if not (path.exists() and path.read_bytes() == blob):
        path.write_bytes(blob)
