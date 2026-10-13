"""Publication guard and identifier check (spec 10.5).

`guard_public(root)` checks a public tree: identifiers, dealers, links to audi.pl, minimum group size and layout.
`identifier_check(root, public=False)` checks any tree (the private repository's data/ included) for VIN-shaped
strings and keys of discarded fields. Both return a list of findings; empty means pass.
"""

from __future__ import annotations

import gzip
import json
import re
from pathlib import Path
from typing import Any

VIN_RE = re.compile(r"\b[A-HJ-NPR-Z0-9]{17}\b")
AUDI_URL_RE = re.compile(r"https?://(?:[a-z0-9-]+\.)*audi\.pl(?![a-z0-9-])", re.IGNORECASE)
IDENTIFIER_KEYS = {"vin", "komisja", "link", "kod_audi", "telefon", "phone"}
PUBLIC_ONLY_KEYS = {"offer_id"}
DEALER_KEYS = {"dealer", "dealer_nr", "dealers"}
# Price, rate, discount and score statistics: each needs `priced` >= min group in the same object.
STAT_KEYS = {
    "special_min", "special_median", "special_max", "discount_pct_median", "discount_pct_max", "catalog_median",
    "price_30d_median", "rate_min", "rate_median", "initial_fee", "score_best", "score_median", "components",
    "offers_80", "lfl_index", "lowest_version_special", "start_price", "drop", "drop_pct", "hot_score", "old", "new",
    "diff", "pct", "median_drop_pct",
}
# Badge and slider counts: each needs `offers` >= min group in the same object.
COUNT_GROUP_KEYS = {"promo_counts", "slider_counts", "promo_changes"}
PUBLIC_LAYOUT = re.compile(
    r"manifest\.json|snapshots/index\.json|dashboard/(?:latest|trends|deals)\.json|"
    r"dashboard/changes/\d{4}-\d{2}-\d{2}\.json|models/[a-z0-9-]+\.json|meta/models\.json")


def _load(path: Path) -> tuple[str, Any]:
    """Return ('json', obj) or ('text', str) for a file; .gz is read decompressed, .xlsx by cell text."""
    name = path.name
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=False)
        cells = [str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        return "text", "\n".join(cells)
    raw = path.read_bytes()
    if name.endswith(".gz"):
        raw = gzip.decompress(raw)
        name = name[:-3]
    text = raw.decode("utf-8-sig", errors="replace")
    if name.endswith(".json"):
        try:
            return "json", json.loads(text)
        except ValueError:
            return "text", text
    return "text", text


def _ok(value: Any, count: Any, i: int | None, mg: int) -> bool:
    c = count[i] if isinstance(count, list) and i is not None and i < len(count) else count
    if isinstance(c, list):
        c = None
    return isinstance(c, int) and c >= mg


def _check_stat(rel: str, path: str, key: str, value: Any, count: Any, count_name: str, mg: int,
                findings: list[str]) -> None:
    if value is None:
        return
    if count is None:
        findings.append(f"{rel}:{path}.{key}: statistic without its `{count_name}` count")
        return
    if isinstance(value, list):
        for i, v in enumerate(value):
            if v is not None and not _ok(v, count, i, mg):
                findings.append(f"{rel}:{path}.{key}[{i}]: rests on fewer than {mg} {count_name} offers")
                return
    elif isinstance(value, dict) and key in COUNT_GROUP_KEYS:
        for k, v in value.items():
            _check_stat(rel, f"{path}.{key}", k, v, count, count_name, mg, findings)
    elif not _ok(value, count, None, mg):
        findings.append(f"{rel}:{path}.{key}: rests on fewer than {mg} {count_name} offers")


def _walk(rel: str, obj: Any, path: str, mg: int, public: bool, findings: list[str]) -> None:
    if isinstance(obj, dict):
        for k in obj:
            kl = str(k).lower()
            if kl in IDENTIFIER_KEYS or (public and kl in PUBLIC_ONLY_KEYS):
                findings.append(f"{rel}:{path}.{k}: identifier key")
            if public and kl in DEALER_KEYS:
                findings.append(f"{rel}:{path}.{k}: dealer key")
            if VIN_RE.search(str(k)):
                findings.append(f"{rel}:{path}: VIN-shaped key")
        if public and mg > 0:
            for k in STAT_KEYS & obj.keys():
                # a dict under a statistic's name is a wrapper object, checked on its own (except components)
                if isinstance(obj[k], dict) and k != "components":
                    continue
                _check_stat(rel, path, k, obj[k], obj.get("priced"), "priced", mg, findings)
            for k in COUNT_GROUP_KEYS & obj.keys():
                # badge and slider counts are dicts; a list or number under the same name is something else
                if not isinstance(obj[k], dict):
                    continue
                _check_stat(rel, path, k, obj[k], obj.get("offers"), "offers", mg, findings)
                if k == "promo_changes" and "offers_before" in obj:
                    _check_stat(rel, path, k, obj[k], obj.get("offers_before"), "offers", mg, findings)
        for k, v in obj.items():
            _walk(rel, v, f"{path}.{k}", mg, public, findings)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            _walk(rel, v, f"{path}[{i}]", mg, public, findings)
    elif isinstance(obj, str):
        _check_text(rel, obj, path, public, findings)


def _check_text(rel: str, text: str, where: str, public: bool, findings: list[str]) -> None:
    m = VIN_RE.search(text)
    if m:
        findings.append(f"{rel}:{where}: VIN-shaped string")
    if public and AUDI_URL_RE.search(text):
        findings.append(f"{rel}:{where}: link to audi.pl")


def _files(root: Path) -> list[Path]:
    if not root.exists():
        print(f"note: {root} does not exist; nothing to check")
        return []
    return sorted(p for p in root.rglob("*") if p.is_file())


def identifier_check(root: Path, public: bool = False) -> list[str]:
    """VIN-shaped strings and keys of discarded fields anywhere under `root` (spec 10.5, first row)."""
    findings: list[str] = []
    for p in _files(root):
        rel = p.relative_to(root).as_posix()
        kind, obj = _load(p)
        if kind == "json":
            _walk(rel, obj, "$", 0, False, findings)
            if public:
                _walk_keys_only(rel, obj, findings)
        else:
            _check_text(rel, obj, "text", False, findings)
    return findings


def _walk_keys_only(rel: str, obj: Any, findings: list[str]) -> None:
    if isinstance(obj, dict):
        for k, v in obj.items():
            if str(k).lower() in PUBLIC_ONLY_KEYS:
                findings.append(f"{rel}: identifier key {k}")
            _walk_keys_only(rel, v, findings)
    elif isinstance(obj, list):
        for v in obj:
            _walk_keys_only(rel, v, findings)


def guard_public(root: Path, min_group: int = 3) -> list[str]:
    """All checks of spec 10.5 on a public tree."""
    findings: list[str] = []
    for p in _files(root):
        rel = p.relative_to(root).as_posix()
        if not PUBLIC_LAYOUT.fullmatch(rel):
            findings.append(f"{rel}: file outside the public layout")
            continue
        kind, obj = _load(p)
        if kind == "json":
            _walk(rel, obj, "$", min_group, True, findings)
        else:
            _check_text(rel, obj, "text", True, findings)
    return findings
